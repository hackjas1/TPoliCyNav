"""
Generates the official 12-slide Final Presentation Deck for TPoliCynav Defense
in TPC Institutional Green Colors using python-pptx.
"""

import pptx
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.enum.shapes import MSO_SHAPE

prs = pptx.Presentation()
prs.slide_width = Inches(13.333)
prs.slide_height = Inches(7.5)

# Color Palette: TPC Institutional Colors
DARK_GREEN = RGBColor(6, 78, 59)      # #064E3B
EMERALD = RGBColor(16, 185, 129)      # #10B981
LIGHT_GREEN = RGBColor(209, 250, 229) # #D1FAE5
WHITE = RGBColor(255, 255, 255)
TEXT_DARK = RGBColor(15, 23, 42)      # Slate 900
GRAY_BG = RGBColor(240, 253, 244)     # #F0FDF4

slides_data = [
    {
        'title': 'TPoliCynav: The TPC Policy Navigator',
        'subtitle': (
            'An AI-Driven Retrieval-Augmented Generation Conversational Agent for Institutional Policy Navigation\n'
            'Talibon Polytechnic College (TPC) • Bohol, Philippines\n\n'
            'IS Professional Elective #4: Generative AI Systems — Capstone Final Defense\n'
            'Team Kuys: Eduardo Avenido Jr. | Christian Paul Larosa | Nielbert Garcia | Marcel Avergonzado\n'
            'Instructor: Jessie A. Melendres | 1st Sem, A.Y. 2026-2027'
        ),
        'is_title': True
    },
    {
        'title': '1. Problem Statement & Motivation',
        'points': [
            'Cognitive Friction in Policy Navigation: Over 4,900 TPC students must navigate a 202-page Student Handbook, municipal ordinances, and scholarship policies.',
            'Fragmented Regulations: Policies governing freshmen entrance attire, boarding house curfews (Ordinance 127), and CHED/GSIS subsidies are isolated across separate documents.',
            'Limitations of Keyword Search: Traditional Ctrl+F keyword search fails to handle paraphrasing, synonymy, or multi-document institutional reasoning.',
            'Target Beneficiaries: Enrolled students, freshmen examinees, boarding house residents, and campus discipline officers (OSAS).'
        ]
    },
    {
        'title': '2. End-to-End System Architecture',
        'points': [
            'End-to-End RAG Pipeline: Raw Corpus -> Text Preprocessing -> BGE-Small Embeddings -> Pinecone Cloud -> LangChain -> Streamlit Web UI -> Railway Cloud.',
            'Embedding Space: BAAI/bge-small-en-v1.5 (384-dimensional dense vectors, MTEB leader, 133 MB footprint).',
            'Vector Database: Pinecone Serverless Index (AWS us-east-1, Cosine metric) paired with a local NumPy matrix fallback.',
            'Orchestration & LLM: LangChain conversational chain paired with Qwen2.5-72B-Instruct / Local Precision Extractor.',
            'Deployment & Secrets: Streamlit Web UI containerized via Multi-Stage Dockerfile, deployed on Railway Cloud via GitHub CI/CD.'
        ]
    },
    {
        'title': '3. Foundational Data Pipeline & Ingestion',
        'points': [
            'Comprehensive Corpus: 7 official collegiate/municipal sources (>260 KB text): Handbook v.04, Ordinance 127, Uniform Guide, Admissions, CHED RA 10931, GSIS STEM, PNP Memorandums.',
            'Automated Ingestion Pipeline (scripts/ingest_pipeline.py): Command-line pipeline parsing PDF/TXT files end-to-end with minimal manual steps.',
            'Sanitization & Normalization: Strips OCR noise, form feeds, page numbers, and standardizes Unicode characters.',
            'Sliding-Window Semantic Chunking: 400-word chunks with 80-word overlap (20% overlap), producing 128 metadata-rich chunks with SHA-256 hashes.'
        ]
    },
    {
        'title': '4. Vector Embeddings & Distance Metric Benchmark',
        'points': [
            'Model Evaluation: BAAI/bge-small-en-v1.5 chosen for exceptional retrieval efficiency and local CPU execution.',
            'Metric Comparison: Evaluated Cosine Similarity, Euclidean Distance (L2), and Dot Product.',
            'Mathematical Equivalence: On unit-normalized vectors (||u||=1), L2^2 = 2 - 2 * Cosine. All three metrics yield identical top-k rankings.',
            'Metric Selection Rationale: Cosine Similarity selected for its bounded [-1, 1] range, enabling an intuitive confidence cutoff (>= 0.50) for out-of-scope refusal.'
        ]
    },
    {
        'title': '5. Prompt Engineering Architecture & Guardrails',
        'points': [
            '3-Tier Structured Prompts: (1) Strict Grounded Policy Assistant, (2) Student Conduct Advisor, (3) Scholarship Evaluator.',
            'Negative Constraints: Explicitly instructs model to state lack of information and redirect to OSAS if context is insufficient.',
            'Pre-Generation Confidence Gate: Maximum cosine score < 0.50 triggers immediate refusal without calling LLM.',
            'Zero Hallucination Verified: Non-campus inquiries (e.g. Tokyo population) score 42.89% and are 100% intercepted.'
        ]
    },
    {
        'title': '6. LangChain Orchestration & Conversational Memory',
        'points': [
            'LangChain Integration: Custom TPCVectorRetriever implementing BaseRetriever interface.',
            'Conversational Memory (app/memory_manager.py): 5-turn sliding window tracking dialogue history.',
            'Coreference Resolution: Contextualizer reformulates follow-up queries (e.g. "Can I wear them?" -> "Can I wear ripped jeans during exams?").',
            'Hybrid Fallback: Gracefully switches between Cloud (Pinecone + Qwen-72B) and Offline (NumPy + Local Extractor).'
        ]
    },
    {
        'title': '7. Application Interface & Institutional Branding',
        'points': [
            'Official TPC Green Theme: Styled with Dark Green (#064E3B), Emerald (#10B981), and Crisp White (#FFFFFF).',
            'Zero-Login Public Access: Deployed web application accessible instantly from any browser without installation.',
            'Frequent Inquiries Quick-Chips: One-click buttons for Degree Programs, Dress Code, Curfew, and GSIS Scholarships.',
            'Transparent Citations: Expandable drawers displaying verbatim policy excerpts and confidence match percentages.'
        ]
    },
    {
        'title': '8. Model Customization: PEFT / LoRA / QLoRA Rationale',
        'points': [
            'Mathematical Mechanics: LoRA updates Delta W = (alpha/r) * B * A, reducing trainable weights by >99.8%.',
            'Hardware Feasibility: QLoRA (NF4 quantization) allows 7B fine-tuning on a single 16 GB VRAM GPU.',
            'Why RAG was Chosen over Fine-Tuning: Policies update frequently (RAG updates in seconds; fine-tuning requires retraining); RAG guarantees auditable citations; zero catastrophic forgetting.',
            'Future LoRA Adaptation: Documented experimental design for Mistral-7B to adapt colloquial Cebuano-English student phrasing.'
        ]
    },
    {
        'title': '9. Containerization, CI/CD & Railway Cloud',
        'points': [
            'Multi-Stage Dockerfile: Stage 1 (Builder: compiles wheels) + Stage 2 (Runner: non-root appuser, >45% size reduction).',
            'Secrets Protocol: PINECONE_API_KEY and HF_TOKEN injected dynamically via Railway dashboard; zero secrets in Git.',
            'Continuous Integration: GitHub Actions workflow (.github/workflows/ci.yml) running automated linting and benchmark tests.',
            'Public Deployment: Hosted live on Railway Cloud via GitHub integration with automatic SSL and zero-install access.'
        ]
    },
    {
        'title': '10. Empirical Benchmark & Quantitative Evaluation',
        'points': [
            'Evaluation Suite (scripts/demo_evaluation.py): 6 structured queries across all institutional domains.',
            'Overall Benchmark Pass Rate: 6/6 Tests Passed (100.0% Pass Rate).',
            'In-Scope Match Confidence: Average 77.26% (Max 85.94% on GSIS subsidy query).',
            'System Latency: Average response time of 2,579 ms (2.58 seconds) on cloud infrastructure.'
        ]
    },
    {
        'title': '11. Limitations, Future Work & Conclusion',
        'points': [
            'Demonstration of All 4 CILOs: End-to-end data pipeline, scalable vector RAG, PEFT rationale, containerized CI/CD deployment.',
            'Limitations: Handling of complex tabular GWA retention charts; informal Cebuano colloquialisms.',
            'Future Roadmap: Bilingual LoRA adapter; voice-enabled campus kiosk mode.',
            'Open Q&A / Defense: Thank you! Team Kuys welcomes questions from the panel.'
        ]
    }
]

blank_slide_layout = prs.slide_layouts[6]

for slide_idx, sdata in enumerate(slides_data):
    slide = prs.slides.add_slide(blank_slide_layout)

    # Background shape
    bg = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, Inches(13.333), Inches(7.5))
    bg.fill.solid()
    bg.fill.fore_color.rgb = DARK_GREEN if sdata.get('is_title') else GRAY_BG
    bg.line.fill.background()

    if sdata.get('is_title'):
        txBox = slide.shapes.add_textbox(Inches(1.0), Inches(1.5), Inches(11.333), Inches(4.5))
        tf = txBox.text_frame
        tf.word_wrap = True

        p = tf.paragraphs[0]
        p.text = 'TALIBON POLYTECHNIC COLLEGE'
        p.font.size = Pt(20)
        p.font.bold = True
        p.font.color.rgb = LIGHT_GREEN
        p.alignment = PP_ALIGN.CENTER

        p2 = tf.add_paragraph()
        p2.text = sdata['title']
        p2.font.size = Pt(40)
        p2.font.bold = True
        p2.font.color.rgb = WHITE
        p2.alignment = PP_ALIGN.CENTER

        p3 = tf.add_paragraph()
        p3.text = '\n' + sdata['subtitle']
        p3.font.size = Pt(16)
        p3.font.color.rgb = LIGHT_GREEN
        p3.alignment = PP_ALIGN.CENTER
    else:
        # Header banner
        header = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, Inches(13.333), Inches(1.3))
        header.fill.solid()
        header.fill.fore_color.rgb = DARK_GREEN
        header.line.fill.background()

        txBox = slide.shapes.add_textbox(Inches(0.8), Inches(0.2), Inches(11.733), Inches(0.9))
        tf = txBox.text_frame
        p = tf.paragraphs[0]
        p.text = sdata['title']
        p.font.size = Pt(28)
        p.font.bold = True
        p.font.color.rgb = WHITE

        # Content box
        contentBox = slide.shapes.add_textbox(Inches(0.8), Inches(1.6), Inches(11.733), Inches(5.3))
        ctf = contentBox.text_frame
        ctf.word_wrap = True

        for i, pt in enumerate(sdata['points']):
            p_pt = ctf.paragraphs[0] if i == 0 else ctf.add_paragraph()
            p_pt.text = '• ' + pt
            p_pt.font.size = Pt(18)
            p_pt.font.color.rgb = TEXT_DARK
            p_pt.space_after = Pt(16)

prs.save('docs/TPoliCynav_Final_Presentation.pptx')
print('Successfully generated docs/TPoliCynav_Final_Presentation.pptx')

